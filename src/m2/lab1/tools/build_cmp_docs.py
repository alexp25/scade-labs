"""Generate the Module 2 Lab 1 Configuration Management Plan (CMP) documents.

One content definition produces two Word files:

  docs/m2/lab1/CMP_template_CruiseControl.docx          (student template, PUBLISHED)
  src/m2/lab1/solution/CMP_CruiseControl_reference.docx (instructor reference, NOT published)

Every piece of content is either plain text (identical in both documents) or an
R(reference, template) pair: the reference document gets the full answer, the
template gets a partial text / "TODO" prompt for the student to complete.

The section structure follows the CMP content areas of IEEE Std 828-2012,
Annex D (D.1-D.9). Text is paraphrased, not copied from the standard.

Usage (from the repo root):
    python src/m2/lab1/tools/build_cmp_docs.py
"""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

REPO = Path(__file__).resolve().parents[4]
TEMPLATE_OUT = REPO / "docs" / "m2" / "lab1" / "CMP_template_CruiseControl.docx"
REFERENCE_OUT = REPO / "src" / "m2" / "lab1" / "solution" / "CMP_CruiseControl_reference.docx"

TODO_COLOR = RGBColor(0xC6, 0x28, 0x28)


class R:
    """Reference text vs. template text for one paragraph or table cell."""

    def __init__(self, reference, template):
        self.reference = reference
        self.template = template


def pick(item, variant):
    return getattr(item, variant) if isinstance(item, R) else item


# ---------------------------------------------------------------------------
# Content definition
# Each section: (heading, [blocks]); a block is ("p", text) / ("bullets", [..])
# / ("table", header, rows) / ("h2", text).
# ---------------------------------------------------------------------------

CI_HEADER = ["CI ID", "Configuration item", "Repository path", "Type", "Baseline(s)"]
CI_ROWS = [
    ["CC-REQ-001", "Cruise Control requirements (REQ-01, 02, 04, 07, 08)", "requirements/requirements.md", "Source", "REQ, DES, VER, REL"],
    ["CC-MDL-001", "Scade One project file", "CruiseControl/CruiseControl.sproj", "Source", "DES, VER, REL"],
    ["CC-MDL-002", "Cruise control design (state machine, regulator, limiter)", "CruiseControl/assets/CC_design.swan", "Source", "DES, VER, REL"],
    [R("CC-MDL-003", "TODO"), R("Vehicle plant model", "TODO"), "CruiseControl/assets/Car_design.swan", R("Source", "TODO"), R("DES, VER, REL", "TODO")],
    [R("CC-MDL-004", "TODO"), R("Closed-loop / open-loop simulation wiring", "TODO"), "CruiseControl/assets/Simulation.swan", R("Source", "TODO"), R("DES, VER, REL", "TODO")],
    [R("CC-TST-001", "TODO"), R("Test scenarios TC-01 ... TC-06", "TODO"), "CruiseControl/scenarios/*.csv", R("Source", "TODO"), R("VER, REL", "TODO")],
    [R("CC-TST-002", "TODO"), R("Evaluation script (scenario runner + report)", "TODO"), "CruiseControl/evaluate_cc_full_report.py", R("Source", "TODO"), R("VER, REL", "TODO")],
    [R("CC-GEN-001", "TODO"), R("Generated Python wrapper + DLL (cc_wrapper)", "TODO"), R("CruiseControl/cc_wrapper/", "TODO"), R("Generated (do not edit by hand)", "TODO"), R("REL", "TODO")],
    [R("CC-TOOL-001", "TODO"), R("Tool configuration: requirements.txt, generate_python_wrapper.bat, Scade One version", "TODO"), R("CruiseControl/requirements.txt, CruiseControl/generate_python_wrapper.bat", "TODO"), R("Source", "TODO"), R("VER, REL", "TODO")],
    [R("CC-VER-001", "TODO"), R("Verification results: results/summary.csv + plots", "TODO"), R("verification/", "TODO"), R("Record (output of a controlled run)", "TODO"), R("VER, REL", "TODO")],
    [R("CC-DOC-001", "TODO"), R("This Configuration Management Plan", "TODO"), R("docs/CMP_CruiseControl.docx", "TODO"), R("Source", "TODO"), R("REQ, DES, VER, REL", "TODO")],
]

BASELINE_HEADER = ["Baseline", "Establishing event", "Controlled items", "Approval authority", "Git tag"]
BASELINE_ROWS = [
    ["Requirements (REQ)", "Requirements reviewed and accepted", "CC-REQ-*", "Project lead", "baseline/req-v1.0"],
    [R("Design (DES)", "TODO"), R("Model builds and simulates without errors; design review passed", "TODO"), R("CC-REQ-*, CC-MDL-*", "TODO"), R("Project lead + reviewer", "TODO"), R("baseline/des-v1.0", "TODO")],
    [R("Verification (VER)", "TODO"), R("All scenarios executed, summary.csv all PASS", "TODO"), R("CC-REQ-*, CC-MDL-*, CC-TST-*, CC-TOOL-*, CC-VER-*", "TODO"), R("Verification engineer", "TODO"), R("baseline/ver-v1.0", "TODO")],
    [R("Release (REL)", "TODO"), R("Verification baseline approved, generated code rebuilt from the tagged model", "TODO"), R("All CIs", "TODO"), R("CM manager (CCB)", "TODO"), R("v1.0.0 (+ GitHub Release)", "TODO")],
]

BRANCH_HEADER = ["Branch", "Purpose", "Created from", "Merges into", "Protection"]
BRANCH_ROWS = [
    ["main", "Approved configuration only; every commit is a released or baselined state", "-", "-", "No direct push; merge via reviewed pull request only"],
    ["develop", "Integration of completed, reviewed work", "main", "release/x.y", R("Merge via pull request", "TODO")],
    [R("feature/REQ-<id>-<short-name>", "TODO"), R("One change request / one requirement at a time", "TODO"), R("develop", "TODO"), R("develop", "TODO"), R("None; deleted after merge", "TODO")],
    [R("release/<major>.<minor>", "TODO"), R("Stabilisation for delivery: only fixes, no new features", "TODO"), R("develop", "TODO"), R("main (then tagged) and back into develop", "TODO"), R("Merge via pull request", "TODO")],
]

RESP_HEADER = ["Role", "CM responsibilities"]
RESP_ROWS = [
    ["CM manager", "Owns this plan, creates baseline and release tags, chairs the change control board (CCB)"],
    [R("Developer", "TODO"), R("Works in feature branches, writes REQ-tagged commit messages, opens pull requests", "TODO")],
    [R("Reviewer", "TODO"), R("Reviews every pull request (model + tests + requirements together) before merge", "TODO")],
    [R("Verification engineer", "TODO"), R("Runs the scenario suite on the candidate configuration and records the results as a CI", "TODO")],
]

SCHEDULE_HEADER = ["Milestone", "CM activity", "Target"]
SCHEDULE_ROWS = [
    ["M1 - Repository created", "Repository, branch protection and this CMP committed", "Week 1"],
    [R("M2 - Requirements baseline", "TODO"), R("Tag baseline/req-v1.0", "TODO"), R("Week 1", "TODO")],
    [R("M3 - Design baseline", "TODO"), R("Tag baseline/des-v1.0 after design review", "TODO"), R("Week 2", "TODO")],
    [R("M4 - Verification baseline", "TODO"), R("Tag baseline/ver-v1.0; configuration audit of the tagged state", "TODO"), R("Week 3", "TODO")],
    [R("M5 - Release", "TODO"), R("Tag v1.0.0, publish GitHub Release with release notes", "TODO"), R("Week 3", "TODO")],
]

SECTIONS = [
    ("1. Introduction (IEEE 828 D.1)", [
        ("h2", "1.1 Purpose of the plan"),
        ("p", "This plan defines how the configuration of the Cruise Control project (built in Lab 4 with Ansys Scade One) "
              "is identified, controlled, recorded and released. It is written for everyone who changes, reviews, "
              "verifies or releases any project artifact."),
        ("h2", "1.2 Scope"),
        ("p", R("CM applies to every artifact that can influence the behavior, the verification or the safety argument of "
                "the Cruise Control software: requirements, Scade One / Swan models, tests, generated code, tool "
                "configuration, verification results and this plan. It covers the project from the requirements "
                "baseline to the first release (v1.0.0). Assumption: a single public GitHub repository is the only "
                "controlled repository.",
                "CM applies to every artifact that can influence the behavior, the verification or the safety argument of "
                "the Cruise Control software. TODO: state which life-cycle phases this plan covers, and list the "
                "assumptions it is based on.")),
        ("h2", "1.3 Relationship to the organization and other projects"),
        ("p", R("Educational project, one team. Requirements come from Lab 3 (REQ-01, REQ-02, REQ-04, REQ-07, REQ-08); the "
                "model and tests come from Lab 4. Deliverable at project close: release v1.0.0 with its release notes.",
                "TODO: describe the product, where its inputs come from (Labs 3 and 4), and what is delivered at project close.")),
        ("h2", "1.4 Key terms"),
        ("bullets", [
            "Configuration item (CI): an artifact placed under configuration control and uniquely identified.",
            "Baseline: an identified, approved configuration used as a reference point; changed only through change control.",
            "Change request (CR): a GitHub issue describing a proposed change and the requirement(s) it affects.",
            R("Release: a baseline delivered outside the team, identified by a semantic version tag (vMAJOR.MINOR.PATCH).",
              "TODO: define 'release'."),
        ]),
        ("h2", "1.5 References"),
        ("bullets", [
            "IEEE Std 828-2012, IEEE Standard for Configuration Management in Systems and Software Engineering.",
            "Module 2, Lesson 1 slides: Configuration Management.",
            "Lab 3 (requirements) and Lab 4 (Cruise Control model) instructions.",
        ]),
    ]),
    ("2. Criteria for identifying configuration items (IEEE 828 D.2)", [
        ("p", "Rule: if an artifact can influence the behavior, the verification or the safety argument of the software, "
              "its version is controlled. Local, user-specific or reproducible-by-rebuild files are excluded."),
        ("p", R("Excluded from CM (listed in .gitignore): jobs/ (Scade One code-generation job folders), *.sproj.user, "
                "*.sproj.editor, pyscadeone.log, __pycache__/, and the working output folder "
                "CruiseControl/results/ (the run that supports a baseline is copied into verification/ instead).",
                "TODO: list the files and folders that are NOT configuration items, and why. "
                "Hint: look at what Scade One creates on your machine that is user-specific or a by-product of a build.")),
    ]),
    ("3. Limitations and assumptions (IEEE 828 D.3)", [
        ("bullets", [
            R("The team uses Scade One Student Edition; the exact tool version is recorded in README.md.",
              "TODO: which tool version assumption must be recorded, and where?"),
            R("GitHub is available; no other repository or backup is in scope.",
              "TODO: add at least one more assumption or limitation."),
        ]),
    ]),
    ("4. CM responsibilities and authorities (IEEE 828 D.4)", [
        ("p", "Changes to a baseline are approved by the change control board (CCB). In this project the CCB is the "
              "CM manager plus one reviewer; approval is recorded as an approved GitHub pull request review."),
    ]),
    ("5. Project organization (IEEE 828 D.5)", [
        ("p", R("One student team. Organizational units: development, review, verification. All interact through the "
                "GitHub repository (issues, pull requests, tags).",
                "TODO: list the organizational units (or people) and how they interact.")),
    ]),
    ("6. CM responsibilities (IEEE 828 D.6)", [
        ("table", RESP_HEADER, RESP_ROWS),
    ]),
    ("7. Applicable policies, directives and procedures (IEEE 828 D.7)", [
        ("bullets", [
            "No direct commits to main; every change reaches main through a reviewed pull request.",
            R("Every commit message starts with the requirement ID(s) or issue it addresses, e.g. 'REQ-07: tune regulator gain (#3)'.",
              "TODO: define the commit message convention."),
            R("Generated code is never edited by hand; it is regenerated from the tagged model.",
              "TODO: what is the policy for generated code?"),
        ]),
    ]),
    ("8. Planned CM activities, schedule and resources (IEEE 828 D.8)", [
        ("h2", "8.1 Configuration identification - configuration items"),
        ("p", "Naming convention for CI IDs: CC-<category>-<nnn>, with categories REQ (requirements), MDL (models), "
              "TST (tests), GEN (generated code), TOOL (tool configuration), VER (verification results), DOC (documents)."),
        ("table", CI_HEADER, CI_ROWS),
        ("h2", "8.2 Controlled repository and branching strategy"),
        ("p", "Controlled repository: one public GitHub repository. Folder layout: requirements/ (requirements), "
              "CruiseControl/ (the Lab 4 Scade One project folder, layout unchanged because its scripts rely on it: "
              "model, scenarios, evaluation script, generated wrapper, tool configuration), verification/ (results of "
              "the run used for the verification baseline) and docs/ (this plan)."),
        ("table", BRANCH_HEADER, BRANCH_ROWS),
        ("h2", "8.3 Versioning"),
        ("p", "Releases use semantic versions vMAJOR.MINOR.PATCH (IEEE 828 Annex K build-naming idea):"),
        ("bullets", [
            "MAJOR - significant change in functionality or interface (e.g. a new input on the cruise_control node).",
            R("MINOR - smaller functional change or important fix (e.g. a new requirement implemented).",
              "MINOR - TODO: when is the minor number incremented?"),
            R("PATCH - minor defect fix with no interface change (e.g. a corrected limiter bound).",
              "PATCH - TODO: when is the patch number incremented?"),
            R("Release candidates carry the suffix -rcN (v1.0.0-rc1, v1.0.0-rc2); nothing may change between the last rc and the release.",
              "TODO: how are release candidates named, and what rule applies between the last candidate and the release?"),
        ]),
        ("p", R("Baseline tags use the form baseline/<type>-v<major>.<minor>, e.g. baseline/req-v1.0, baseline/req-v1.1.",
                "TODO: define the naming form of baseline tags.")),
        ("h2", "8.4 Baselines"),
        ("table", BASELINE_HEADER, BASELINE_ROWS),
        ("p", R("Release manifest example: v1.0.0 = baseline/req-v1.0 + baseline/des-v1.1 + baseline/ver-v1.0 + "
                "Scade One version + requirements.txt (tool configuration) + CruiseControl/cc_wrapper rebuilt from the tagged model.",
                "TODO: write the manifest of your first release (which baseline of each kind it contains).")),
        ("h2", "8.5 Change control"),
        ("bullets", [
            "1. Open a GitHub issue (change request) naming the affected REQ ID(s).",
            "2. Create feature/REQ-<id>-<short-name> from develop.",
            R("3. Update requirements, model and tests together; commit with REQ-tagged messages.",
              "3. TODO"),
            R("4. Open a pull request referencing the issue; a reviewer checks model + tests + requirements and re-runs the scenarios.",
              "4. TODO"),
            R("5. Merge into develop; the issue is closed by the merge.", "5. TODO"),
            R("6. If a baseline is affected, the CM manager creates the next baseline tag.", "6. TODO"),
        ]),
        ("h2", "8.6 Configuration status accounting"),
        ("p", R("Status is reported from GitHub: open/closed issues per REQ ID, merged pull requests, the list of tags "
                "and releases. The release notes of each GitHub Release list the issues and requirements included.",
                "TODO: how will the team know, at any time, which change is in which version?")),
        ("h2", "8.7 Configuration audits"),
        ("p", R("Before each release: functional audit (every requirement in the baseline has a passing scenario in "
                "verification/summary.csv) and physical audit (a fresh clone of the tag rebuilds the wrapper and "
                "reproduces the same results).",
                "TODO: describe the functional and physical audit you will perform before release.")),
        ("h2", "8.8 CM schedule"),
        ("table", SCHEDULE_HEADER, SCHEDULE_ROWS),
        ("h2", "8.9 CM resources"),
        ("bullets", [
            "Git, a GitHub account and a public repository.",
            R("Ansys Scade One Student Edition and Python 3.12 with ansys-scadeone-core (for regenerating the wrapper).",
              "TODO: list the tools needed to rebuild and verify the configuration."),
        ]),
    ]),
    ("9. CMP maintenance (IEEE 828 D.9)", [
        ("p", R("The CM manager reviews this plan at each baseline. Changes to the plan follow the same change control "
                "process as any other CI (it is CI CC-DOC-001); the document history table below records each revision.",
                "TODO: who maintains this plan, how often, and how are its changes approved and communicated?")),
    ]),
]

HISTORY_HEADER = ["Version", "Date", "Author", "Change"]


# ---------------------------------------------------------------------------
# Rendering
# ---------------------------------------------------------------------------

def add_text(paragraph, text):
    """Add text to a paragraph, highlighting TODO prompts in red."""
    if "TODO" in text:
        before, _, after = text.partition("TODO")
        if before:
            paragraph.add_run(before)
        todo = paragraph.add_run("TODO" + after)
        todo.font.color.rgb = TODO_COLOR
        todo.italic = True
    else:
        paragraph.add_run(text)


def shade(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def add_table(doc, header, rows, variant):
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    for i, title in enumerate(header):
        cell = table.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(title)
        run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shade(cell, "0D1B3E")
    for row in rows:
        cells = table.add_row().cells
        for i, item in enumerate(row):
            cells[i].text = ""
            add_text(cells[i].paragraphs[0], pick(item, variant))
    doc.add_paragraph()


def build(variant, out_path):
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    title = doc.add_heading("Configuration Management Plan", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run("Cruise Control System (Ansys Scade One)").bold = True
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(meta, pick(R("Document ID: CC-DOC-001  |  Team: reference solution", "Document ID: CC-DOC-001  |  Team: TODO"), variant))

    if variant == "template":
        note = doc.add_paragraph()
        add_text(note, "TODO: This is a partially completed template. Enrich the given content where you can, "
                       "replace every red TODO with your own decisions, and keep the section structure "
                       "(IEEE 828-2012, Annex D). Commit the finished plan as docs/CMP_CruiseControl.docx in your repository.")

    doc.add_heading("Document history", level=1)
    history = [["0.1", "TODO", "TODO", "Initial draft from the course template"]] if variant == "template" else \
        [["0.1", "-", "Instructor", "Initial draft"], ["1.0", "-", "Instructor", "Reference solution"]]
    add_table(doc, HISTORY_HEADER, history, variant)

    for heading, blocks in SECTIONS:
        doc.add_heading(heading, level=1)
        for block in blocks:
            kind = block[0]
            if kind == "h2":
                doc.add_heading(block[1], level=2)
            elif kind == "p":
                add_text(doc.add_paragraph(), pick(block[1], variant))
            elif kind == "bullets":
                for item in block[1]:
                    add_text(doc.add_paragraph(style="List Bullet"), pick(item, variant))
            elif kind == "table":
                add_table(doc, block[1], block[2], variant)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    print(f"wrote {out_path.relative_to(REPO)}")


if __name__ == "__main__":
    build("template", TEMPLATE_OUT)
    build("reference", REFERENCE_OUT)
